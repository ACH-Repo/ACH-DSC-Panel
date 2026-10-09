"""The swipe on selected curves scales their own y axis, so the mass and
the heat flow are made taller or flatter one at a time; and a DTG's peak
temperature.
"""

import json
import os

import pytest

from dscpanel.core import measure, model, session

from test_family import _swipe, clipboard  # noqa: F401 (a fixture)
from test_weight import _window, sdt_data


def _sdt_window(qapp):
    """An SDT run with its mass AND its heat flow shown: two y axes."""
    win = _window(qapp, model.Sample("C:/nowhere/SDT-1.tri", sdt_data()))
    win.toggle_signal(model.SIGNAL_HEAT)
    win.undo.clear()
    win.plot.grab()
    return win


def _curves(doc):
    mass = [s for s in doc.scans if s.is_mass][0]
    heat = [s for s in doc.scans if s.is_heat][0]
    return mass, heat


def _place(plot, scan):
    """Where a curve's baseline is on the page."""
    return plot.sy_to_px(scan, plot._level_of(scan) + scan.offset)


# ------------------------------------------------- the swipe, axis by axis
def test_nothing_selected_swipes_the_main_axis_as_before(qapp):
    win = _sdt_window(qapp)
    doc, plot = win.doc, win.plot
    assert plot.main_axis() == "y2"
    y, y2 = plot.view_y(), plot.view_y2()
    doc.select_all(False)
    assert plot.swiped_axes() == ["y2"]
    _swipe(plot, 2)                                    # 1.2 ** 2 taller
    plot.commit_view()
    assert plot.view_y2() == pytest.approx((y2[0] / 1.44, y2[1] / 1.44))
    assert plot.view_y() == pytest.approx(y)


def test_a_selected_heat_flow_swipes_its_own_axis(qapp):
    """The heat flow taller, the m% left exactly as it was - and the heat
    flow in its place; one undo step puts it back."""
    win = _sdt_window(qapp)
    doc, plot = win.doc, win.plot
    mass, heat = _curves(doc)
    y, y2 = plot.view_y(), plot.view_y2()
    stored = (mass.offset, heat.offset)
    place = _place(plot, heat)
    doc.select_only([heat])
    assert plot.swiped_axes() == ["y"]
    _swipe(plot, 2)
    plot.commit_view()
    plot.grab()
    assert plot.view_y() == pytest.approx((y[0] / 1.44, y[1] / 1.44))
    assert plot.view_y2() == pytest.approx(y2)
    assert mass.offset == stored[0]
    assert _place(plot, heat) == pytest.approx(place, abs=0.5)
    win.undo_step()
    assert plot.view_y() == pytest.approx(y)
    assert (mass.offset, heat.offset) == pytest.approx(stored)


def test_a_selection_on_both_axes_swipes_both(qapp):
    win = _sdt_window(qapp)
    doc, plot = win.doc, win.plot
    mass, heat = _curves(doc)
    y, y2 = plot.view_y(), plot.view_y2()
    doc.select_only([mass, heat])
    assert sorted(plot.swiped_axes()) == ["y", "y2"]
    _swipe(plot, 1)
    plot.commit_view()
    assert plot.view_y() == pytest.approx((y[0] / 1.2, y[1] / 1.2))
    assert plot.view_y2() == pytest.approx((y2[0] / 1.2, y2[1] / 1.2))


def test_a_hidden_or_axisless_selected_curve_does_not_decide(qapp):
    """A selected curve that is not drawn - hidden, or a heat flow beside a
    DTG, which has no axis - leaves the choice to the rest."""
    win = _sdt_window(qapp)
    doc, plot = win.doc, win.plot
    mass, heat = _curves(doc)
    heat.visible = False
    win.refresh()
    doc.select_only([heat])
    assert plot.swiped_axes() == [plot.main_axis()]
    heat.visible = True
    win.toggle_signal(model.SIGNAL_DTG)
    assert doc.axis_missing(heat)
    doc.select_only([heat])
    assert plot.swiped_axes() == [plot.main_axis()]
    dtg_scan = [s for s in doc.scans if s.is_dtg][0]
    doc.select_only([dtg_scan])
    assert plot.swiped_axes() == ["y"]
    assert plot.axis_name("y") == "DTG"


# ------------------------------------------------------ the DTG's peak
def _dtg_window(qapp):
    win = _window(qapp, model.Sample("C:/nowhere/SDT-1.tri", sdt_data()))
    win.toggle_signal(model.SIGNAL_DTG)
    win.undo.clear()
    win.plot.grab()
    return win, [s for s in win.doc.scans if s.is_dtg][0]


def test_a_dtg_offers_its_peak_temperature(qapp):
    _win, dtg_scan = _dtg_window(qapp)
    assert [m.name for m in measure.models_for(dtg_scan)] == [
        measure.DTG_PEAK]
    # and it is a DTG's alone
    mass = [s for s in _win.doc.scans if s.is_mass][0]
    assert measure.DTG_PEAK not in [m.name for m in measure.models_for(mass)]


def test_the_dtg_peak_is_where_the_mass_changes_fastest(qapp):
    """The run loses a quarter of its mass in one step centred on 350 degC:
    the peak is there, by temperatures typed or by the samples a drag ran
    between, and its label says so."""
    win, dtg_scan = _dtg_window(qapp)
    typed = measure.run(measure.DTG_PEAK, dtg_scan, 200.0, 500.0)
    assert typed.value() == pytest.approx(350.0, abs=2.0)
    assert typed.cursors() == pytest.approx([200.0, 500.0])
    dragged = measure.run(measure.DTG_PEAK, dtg_scan, 200.0, 500.0,
                          span=(100, 300))
    assert dragged.value() == pytest.approx(350.0, abs=2.0)
    assert dragged.summary(win.doc).startswith("*T*_{p} = 35")
    assert not dragged.marks_a_point                 # no tangent lines
    # off the step there is no peak to speak of, but a stretch is still
    # measured: the largest departure from the line between its ends
    flat = measure.run(measure.DTG_PEAK, dtg_scan, 100.0, 250.0)
    assert flat is not None and 100.0 <= flat.value() <= 250.0
    # another curve is never measured as a DTG
    mass = [s for s in win.doc.scans if s.is_mass][0]
    assert measure.dtg_peak(mass, 200.0, 500.0) is None


def test_the_dtg_peak_is_drawn_and_survives_a_session(qapp, tmp_path):
    win, dtg_scan = _dtg_window(qapp)
    win.ask_analysis = lambda x0, x1: measure.DTG_PEAK
    made = win._measure_ready(dtg_scan, 200.0, 500.0, None, span=(100, 300))
    assert made is not None and made in dtg_scan.analysis_objects
    win.plot.grab()
    assert [a for a, _box in win.plot._analysis_boxes] == [made]
    path = str(tmp_path / "dtg.dscpanel")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(session.to_state(win.doc), fh)
    loaded, problems = session.load(
        path, lambda p: model.Sample(p, sdt_data()))
    assert not problems
    (again,) = [s for s in loaded.scans if s.is_dtg]
    (peak,) = again.analysis_objects
    assert peak.model_name == measure.DTG_PEAK
    assert peak.value() == pytest.approx(made.value())
    win.undo_step()
    assert made not in dtg_scan.analysis_objects


# ------------------------------------------------------- mass lines
def _mass_window(qapp):
    """An SDT run with its m% alone, a mass at 350 degC on it - the middle
    of its step, 87.5 %."""
    win = _window(qapp, model.Sample("C:/nowhere/SDT-1.tri", sdt_data()))
    mass = [s for s in win.doc.scans if s.is_mass][0]
    at = measure.run(measure.MASS_AT, mass, 350.0, 350.0)
    win.undo.clear()
    win.refresh()
    win.plot.grab()
    return win, mass, at


def _row_ink(image, y, x0, x1):
    """How many pixels of row `y` between x0 and x1 are not the page."""
    page = image.pixelColor(x0, 2).rgb()
    return sum(1 for x in range(x0, x1)
               if image.pixelColor(x, y).rgb() != page)


def test_a_mass_line_marks_the_m_percent_of_its_point(qapp):
    """The template's `showmass`: a dashed line across the axes at the m%
    of the analysis's point, its value at the left edge - off until
    asked for, and the value is the measurement, without the offset."""
    from PySide6.QtCore import QPointF
    win, mass, at = _mass_window(qapp)
    plot = win.plot
    (trace,) = [t for t in plot.traces if t.scan is mass]
    assert at.has_mass_line and not at.mass_line
    assert plot.mass_line(trace, at) is None
    at.mass_line = True
    win.refresh()
    plot.grab()
    (trace,) = [t for t in plot.traces if t.scan is mass]
    y, value = plot.mass_line(trace, at)
    # the curve at the point the label's arrow points at (a sample apart
    # from the marker's own "at or past" sample on this coarse test run)
    assert value == pytest.approx(float(at.fields["Mass"].split()[0]),
                                  abs=0.5)
    assert plot.mass_line_text(value) == "{:.1f} %".format(value)
    # drawn: a dashed row of ink across the axes where there was none
    rect = plot.plot_rect()
    with_line = plot.grab().toImage()
    k = with_line.width() / float(plot.width())
    row = int(round(plot.to_widget(QPointF(rect.left(), y)).y() * k))
    left = int(plot.to_widget(rect.topLeft()).x() * k) + 40
    right = int(plot.to_widget(rect.bottomRight()).x() * k) - 10
    ink = max(_row_ink(with_line, r, left, right)
              for r in (row - 1, row, row + 1))
    assert ink > (right - left) // 4
    at.mass_line = False
    win.refresh()
    without = plot.grab().toImage()
    bare = max(_row_ink(without, r, left, right)
               for r in (row - 1, row, row + 1))
    assert bare < ink // 3
    # moved up the stack, the line goes with the curve, the value stays
    at.mass_line = True
    mass.offset = 5.0
    win.refresh()
    plot.grab()
    (trace,) = [t for t in plot.traces if t.scan is mass]
    y2, value2 = plot.mass_line(trace, at)
    assert value2 == pytest.approx(value)
    assert y2 < y
    # the house style says how it is written
    win.doc.style.mass_line_format = "%.0f"
    assert plot.mass_line_text(87.46) == "87 %"


def test_only_a_point_on_a_mass_curve_has_a_mass_line(qapp):
    win = _sdt_window(qapp)
    mass, heat = _curves(win.doc)
    onset = measure.run("Onset point", mass, 300.0, 400.0)
    assert onset is not None and onset.has_mass_line
    peak = measure.run("Peak height", heat, 100.0, 300.0)
    assert peak is not None and not peak.has_mass_line
    from dscpanel.ui.dialogs import AnalysisSettings
    shown = AnalysisSettings(win, onset, on_change=win._live_change)
    assert not shown.mass_line.isHidden()
    shown.mass_line.setChecked(True)
    assert onset.mass_line
    shown.close()
    hidden = AnalysisSettings(win, peak, on_change=win._live_change)
    assert hidden.mass_line.isHidden()
    hidden.close()


def test_a_mass_line_survives_a_session(qapp, tmp_path):
    win, mass, at = _mass_window(qapp)
    at.mass_line = True
    path = str(tmp_path / "mass.dscpanel")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(session.to_state(win.doc), fh)
    loaded, problems = session.load(
        path, lambda p: model.Sample(p, sdt_data()))
    assert not problems
    (again,) = [s for s in loaded.scans if s.is_mass]
    (marker,) = again.analysis_objects
    assert marker.mass_line


def _lined(qapp):
    """`_mass_window` with its mass line on, painted."""
    win, mass, at = _mass_window(qapp)
    at.mass_line = True
    win.refresh()
    win.plot.grab()
    (trace,) = [t for t in win.plot.traces if t.scan is mass]
    return win, mass, at, trace


def test_a_mass_line_is_thinner_than_the_curves(qapp):
    from dscpanel.core import style
    win, mass, _at, _trace = _lined(qapp)
    curve = win.plot.style_of(mass, "line_width")
    assert style.figure_value(win.doc, "mass_line_width") < curve


def _row_delta(first, second, row, columns):
    """The mean difference of two pictures along `row` (and the rows
    beside it, the line being thin), over `columns`."""
    total = 0.0
    for x in columns:
        total += max(sum(abs(a - b) for a, b in zip(
            first.pixelColor(x, r).getRgb()[:3],
            second.pixelColor(x, r).getRgb()[:3]))
            for r in (row - 1, row, row + 1))
    return total / max(1, len(columns))


def test_a_mass_line_fades_where_it_passes_behind_something(qapp):
    """Under everything, and faint where it passes behind the curve, a
    label or a decorator: what is there is worked out before it is drawn
    (`mass_obstacles`), and the line drawn there at a quarter of its
    colour."""
    from PySide6.QtCore import QPointF
    win, mass, at, trace = _lined(qapp)
    plot = win.plot
    y, _value = plot.mass_line(trace, at)
    rect = plot.plot_rect()
    spans = plot.mass_obstacles(y)
    # its curve crosses it at the point, its own value sits on it
    point_x = float(plot.x_to_px(plot.to_axis(at.value()), rect))
    assert any(a <= point_x <= b for a, b in spans)
    value_box = plot.mass_text_box(trace, at)[0]
    assert any(a <= value_box.center().x() <= b for a, b in spans)
    # a label put on the line is something it passes behind as well
    label = win.add_label("Over the line",
                          at=QPointF(rect.left() + rect.width() * 0.7, y))
    win.refresh()
    plot.grab()
    box = [b for lb, b in plot._text_boxes if lb is label][0]
    spans = plot.mass_obstacles(y)
    assert any(a <= box.center().x() <= b for a, b in spans)
    # drawn faint there: the line adds less to the picture inside a span
    # than along a free stretch
    with_line = plot.grab().toImage()
    at.mass_line = False
    win.refresh()
    without = plot.grab().toImage()
    k = with_line.width() / float(plot.width())
    row = int(round(plot.to_widget(QPointF(rect.left(), y)).y() * k))

    def columns(a, b):
        lo = int(plot.to_widget(QPointF(a, y)).x() * k)
        hi = int(plot.to_widget(QPointF(b, y)).x() * k)
        return list(range(lo, hi))

    lo_box, hi_box = box.left() + 2, box.right() - 2
    free = columns(rect.left() + rect.width() * 0.35,
                   rect.left() + rect.width() * 0.55)
    free = [x for x in free if not any(
        plot.to_widget(QPointF(a, y)).x() * k - 2 <= x
        <= plot.to_widget(QPointF(b, y)).x() * k + 2 for a, b in spans)]
    behind = columns(lo_box, hi_box)
    assert len(free) > 20 and len(behind) > 20
    assert (_row_delta(with_line, without, row, behind)
            < 0.5 * _row_delta(with_line, without, row, free))


def test_the_value_of_a_mass_line_is_dragged_along_it(qapp):
    """Along its line freely, up or down only a little; one undo step.
    A click on it is a click on its analysis."""
    from test_window import _move, _press, _release
    win, mass, at, trace = _lined(qapp)
    plot = win.plot
    box = plot.mass_text_box(trace, at)[0]
    handle = plot.object_at(box.center())
    assert handle is at.mass_text
    start = plot.to_widget(box.center())
    plot.mousePressEvent(_press(plot, (start.x(), start.y())))
    for step in range(1, 6):
        plot.mouseMoveEvent(_move(plot, (start.x() + 30 * step,
                                         start.y() + 40 * step)))
    plot.mouseReleaseEvent(_release(plot, (start.x() + 150,
                                           start.y() + 200)))
    plot.grab()
    moved = plot.mass_text_box(trace, at)[0]
    assert at.mass_at is not None and moved.left() > box.left() + 50
    # far below where the hand went: it stays near its line
    y, _value = plot.mass_line(trace, at)
    assert abs(moved.center().y() - y) <= moved.height() / 2.0 + \
        plot.MASS_REACH + 0.5
    win.undo_step()
    assert at.mass_at is None and at.mass_dy is None
    # a click selects the value itself, not its analysis
    plot.grab()
    win.doc.select_all(False)
    plot.select_at(plot.mass_text_box(trace, at)[0].center())
    assert at.mass_text.selected and not at.selected
    assert win.doc.selected() == [at.mass_text]


def test_g_moves_a_selected_mass_line_value_and_x_locks_it(qapp):
    """Selected, the value is what G moves - not its analysis's label -
    and X keeps it at its height, Y on its place along the line."""
    from PySide6.QtCore import QPointF, Qt
    from PySide6.QtTest import QTest
    win, mass, at, trace = _lined(qapp)
    plot = win.plot
    label_dy = at.label_dy
    win.doc.select_only([at.mass_text])
    plot._cursor = QPointF(plot.plot_rect().center())
    assert plot.start_grab()
    assert plot._move["objs"] == [at.mass_text]
    height = plot.mass_place(at)[1]
    QTest.keyClick(plot, Qt.Key_X)
    start = plot._move["start"]
    plot._update_move(start + QPointF(60.0, 30.0))
    along, dy = plot.mass_place(at)
    assert along > plot.MASS_AT
    assert dy == pytest.approx(height)            # X: not up or down
    QTest.keyClick(plot, Qt.Key_Y)
    plot._update_move(start + QPointF(60.0, 8.0))
    assert plot.mass_place(at)[0] == pytest.approx(plot.MASS_AT)
    QTest.keyClick(plot, Qt.Key_Return)
    assert at.label_dy == label_dy                  # the label never moved
    assert at.mass_dy is not None
    win.undo_step()
    assert at.mass_at is None and at.mass_dy is None
    # H on the selected value takes its line away; one undo brings it back
    win.doc.select_only([at.mass_text])
    win.hide_selected()
    assert not at.mass_line and at.visible
    win.undo_step()
    assert at.mass_line


def test_where_a_mass_line_value_stands_is_saved(qapp, tmp_path):
    win, mass, at, _trace = _lined(qapp)
    at.mass_at, at.mass_dy = 0.4, 6.0
    path = str(tmp_path / "placed.dscpanel")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(session.to_state(win.doc), fh)
    loaded, _problems = session.load(
        path, lambda p: model.Sample(p, sdt_data()))
    (again,) = [s for s in loaded.scans if s.is_mass]
    (marker,) = again.analysis_objects
    assert (marker.mass_at, marker.mass_dy) == (0.4, 6.0)


# ------------------------------------------------ TRIOS's Excel export
DEG = "\u00b0"


def test_sheet_names_are_made_as_trios_makes_them():
    """As a real TRIOS export of a seven-step run names its sheets: no
    "/", "-2" and "-3" after repeats, and the LAST 31 characters."""
    from dscpanel.core import export
    up = "Ramp 10.00 {0}C/min to 250.0000 {0}C".format(DEG)
    down = "Ramp 10.00 {0}C/min to 30.0000 {0}C".format(DEG)
    fast = "Ramp 50.00 {0}C/min to 250.0000 {0}C".format(DEG)
    names = export.trios_sheet_names([up, down, up, down, up, down, fast])
    assert names == [
        "Ramp 10.00 {0}Cmin to 250.0000 {0}C".format(DEG),
        "Ramp 10.00 {0}Cmin to 30.0000 {0}C".format(DEG),
        "mp 10.00 {0}Cmin to 250.0000 {0}C-2".format(DEG),
        "amp 10.00 {0}Cmin to 30.0000 {0}C-2".format(DEG),
        "mp 10.00 {0}Cmin to 250.0000 {0}C-3".format(DEG),
        "amp 10.00 {0}Cmin to 30.0000 {0}C-3".format(DEG),
        "Ramp 50.00 {0}Cmin to 250.0000 {0}C".format(DEG)]
    assert all(len(n) <= 31 for n in names)
    assert export.trios_step_name(up + " #3") == up


def _rows(sheet):
    return list(sheet.iter_rows(values_only=True))


def test_an_sdt_run_is_written_as_trios_writes_it(tmp_path):
    """Details, then a sheet per step: its name, the column names, the
    units, every sample - an SDT run's weight in %, called "Weight"."""
    import numpy as np
    import openpyxl
    from dscpanel.core import export
    data = sdt_data(segments=2)
    data["head"].update({"operator": "An Operator", "rundate": "09/10/2025",
                         "instrumentname": "SDT (bench 1)",
                         "proceduresegments": "Equilibrate 30 C; Ramp"})
    # the instrument's flagged tail: time and weight still measured
    nums = data["numdata"][1]["nums"]
    nums[-4:, 1] = np.nan                               # temperature
    sample = model.Sample("C:/nowhere/SDT-Run-1.tri", data)
    assert len(sample.data["numdata"][1]["tail"]) == 4  # kept aside
    assert export.trios_name(sample) == "sdt-run-1"
    path = str(tmp_path / "sdt-run-1.xlsx")
    export.write_trios_excel(sample, [0, 1], path)
    book = openpyxl.load_workbook(path, read_only=True)
    assert book.sheetnames[0] == "Details"
    assert _rows(book["Details"]) == [
        ("Filename", "sdt-run-1"), ("Instrument name", "SDT (bench 1)"),
        ("Operator", "An Operator"), ("rundate", "09/10/2025"),
        ("Sample name", "SDT-1"),
        ("proceduresegments", "Equilibrate 30 C; Ramp")]
    first, second = book.worksheets[1:]
    rows = _rows(second)
    assert rows[0][0] == "Ramp 10.00 {0}C/min to 650.000 {0}C".format(DEG)
    assert rows[1] == ("Time", "Temperature", "Heat Flow (Normalized)",
                       "Weight")
    assert rows[2] == ("min", DEG + "C", "W/g", "%")
    assert len(rows) == 3 + 400                       # the tail as well
    assert rows[-1][1] is None and rows[-1][3] is not None
    assert rows[3][3] == pytest.approx(float(nums[0, 4]))
    # repeated step names: the second sheet is "...-2"
    assert second.title.endswith("-2") and len(second.title) <= 31


def test_a_dsc_run_without_a_mass_writes_its_heat_flow_in_mw(tmp_path):
    import openpyxl
    from conftest import make_data
    from dscpanel.core import export
    sample = model.Sample("C:/nowhere/TEST-1.tri", make_data())
    path = str(tmp_path / "test-1.xlsx")
    export.write_trios_excel(sample, [0], path)
    rows = _rows(openpyxl.load_workbook(path, read_only=True).worksheets[1])
    assert rows[1] == ("Time", "Temperature", "Heat Flow")
    assert rows[2][2] == "mW"
    watts = sample.data["numdata"][0]["nums"][0, 2]
    assert rows[3][2] == pytest.approx(watts * 1000.0)


def test_the_window_exports_a_workbook_per_file_of_the_figure(
        qapp, tmp_path):
    """The segments on the figure (the selected curves', if any), one
    workbook per file; it asks before replacing one."""
    import openpyxl
    from dscpanel.ui.window import MainWindow
    win = MainWindow()
    first = model.Sample("C:/nowhere/Run-A.tri", sdt_data(segments=2))
    second = model.Sample("C:/nowhere/Run-B.tri", sdt_data())
    win._sample_loaded(first)
    win._sample_loaded(second)
    win.toggle_segment(first, 1, True)
    assert [(s.name, segs) for s, segs in
            export_groups(win)] == [("Run-A", [0, 1]), ("Run-B", [0])]
    written = win.export_excel(str(tmp_path))
    assert sorted(os.path.basename(p) for p in written) == [
        "run-a.xlsx", "run-b.xlsx"]
    book = openpyxl.load_workbook(str(tmp_path / "run-a.xlsx"),
                                  read_only=True)
    assert len(book.sheetnames) == 3                  # Details + 2 steps
    # only the selected curve's segment, and asked before replacing
    win.doc.select_only([s for s in win.doc.scans
                         if s.sample is first and s.seg == 1])
    asked = []
    win.ask_replace = lambda names: asked.append(names) or True
    written = win.export_excel(str(tmp_path / "run-a.xlsx"))
    book = openpyxl.load_workbook(written[0], read_only=True)
    assert len(book.sheetnames) == 2
    assert win.ops.get("file.export_excel") is not None


def export_groups(win):
    from dscpanel.core import export
    return export.trios_excel_groups(win.doc)


def test_a_long_metadata_string_is_read_whole():
    """.NET writes a length of 128 or more in two bytes, seven bits each:
    a procedure of several steps is that long, and was read as missing."""
    from dscpanel.core import trios_io
    text = "Ramp 10.00 C/min to 250.0000 C; " * 5
    length = len(text.encode())
    assert length >= 128
    raw = (b"\x00\x11proceduresegments" + bytes([0x80 | (length & 0x7F),
                                               length >> 7])
           + text.encode() + b"\x00")
    assert trios_io._meta_string(raw, "proceduresegments") == text
    short = b"\x08operator\x04Anna"
    assert trios_io._meta_string(short, "operator") == "Anna"


# ------------------------------------------ the measured data as text
def test_the_measured_data_is_written_as_tab_separated_text(tmp_path):
    """Our own format: what the run records on "#" lines, a header row of
    names with their units, every sample of each segment with its number
    first - an empty sample an empty cell, the tail included."""
    import numpy as np
    from dscpanel.core import export
    data = sdt_data(segments=2)
    data["head"].update({"operator": "An Operator", "rundate": "09/10/2025"})
    nums = data["numdata"][1]["nums"]
    nums[-4:, 1] = np.nan                               # temperature
    sample = model.Sample("C:/nowhere/SDT-Run-1.tri", data)
    assert export.data_text_name(sample) == "SDT-Run-1.csv"
    path = str(tmp_path / "SDT-Run-1.csv")
    export.write_data_text(sample, [0, 1], path)
    with open(path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    comments = [line for line in lines if line.startswith("#")]
    rows = [line.split("\t") for line in lines if not line.startswith("#")]
    assert comments[0] == "# SDT-Run-1.tri, exported by Triplot"
    assert "# Operator: An Operator" in comments
    assert "# Run date: 09/10/2025" in comments
    assert comments[-2].startswith("# Segment 1: ")
    assert comments[-1].startswith("# Segment 2: Ramp")
    assert rows[0] == ["Segment", "Time (min)",
                       "Temperature ({}C)".format(DEG),
                       "Heat Flow (Normalized) (W/g)", "Weight (%)"]
    second = [row for row in rows[1:] if row[0] == "2"]
    assert len(second) == 400                           # the tail as well
    assert second[-1][2] == "" and second[-1][4] != ""
    assert float(second[0][4]) == pytest.approx(float(nums[0, 4]),
                                                rel=1e-7)
    assert len(rows) == 1 + len(second) + sum(
        1 for row in rows[1:] if row[0] == "1")
    assert all(len(row) == 5 for row in rows)


def test_the_window_exports_the_data_a_file_each(qapp, tmp_path):
    """As the Excel export: the segments on the figure, a file each, two
    files of one name told apart."""
    from dscpanel.ui.window import MainWindow
    win = MainWindow()
    first = model.Sample("C:/nowhere/Run-A.tri", sdt_data(segments=2))
    second = model.Sample("C:/elsewhere/Run-A.tri", sdt_data())
    win._sample_loaded(first)
    win._sample_loaded(second)
    written = win.export_data(str(tmp_path))
    assert sorted(os.path.basename(p) for p in written) == [
        "Run-A-2.csv", "Run-A.csv"]
    assert win.ops.get("file.export_data") is not None
    assert "file.export_data" in dict(win.MENUS)["Fi&le"]


# ------------------------------------- analyses pasted onto another figure
def test_analyses_paste_onto_the_same_file_in_another_figure(qapp,
                                                             clipboard):
    """A curve copied in one tab and pasted onto the same file's curve in
    another: TRIOS's own analysis gives its settings to the one of the
    same key there, and one made here is measured again over the SAME
    samples (the same file, segment and length: `session.data_key`)."""
    import numpy as np
    from conftest import make_data
    from dscpanel.ui.window import MainWindow
    from test_family import _choose
    stored = {"Ramp 10,00 C/min to 250 C #1": {
        "Onset point": [{"segment": 1, "Onset x": "80,0 \u00b0C"}]}}
    win = MainWindow()
    win._sample_loaded(model.Sample("C:/nowhere/Same.tri",
                                    make_data(analyses=stored)))
    first = win.doc.scans[0]
    own = first.analysis_objects[0]
    assert own.source != "panel"
    own.visible, own.colour = True, "#00aa55"
    win.plot.grab()
    xs = [t for t in win.plot.traces if t.scan is first][0].x
    made = None
    for entry in measure.models_for(first):
        made = measure.run(entry.name, first, float(xs[100]),
                           float(xs[200]), span=(100, 200))
        if made is not None:
            break
    assert made is not None and made.span == (100, 200)
    win.refresh()
    win.doc.select_only([first])
    win.copy_selected()

    win.new_figure()
    win._sample_loaded(model.Sample("C:/nowhere/Same.tri",
                                    make_data(analyses=stored)))
    second = win.doc.scans[0]
    assert second is not first
    theirs = second.analysis_objects[0]
    assert not theirs.visible
    win.doc.select_only([second])
    _choose(win, "Analyses")
    win.paste()
    assert theirs.visible and theirs.colour == "#00aa55"
    again = [a for a in second.analysis_objects if a.source == "panel"]
    assert len(again) == 1 and again[0].span == (100, 200)
    assert np.isclose(again[0].cursors()[0], made.cursors()[0])
    win.undo.undo()
    assert not theirs.visible
    assert not [a for a in second.analysis_objects if a.source == "panel"]
